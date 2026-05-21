"""Document extraction and text chunking service.

Supports: PDF, DOCX, PPTX, TXT
"""

import os
import fitz  # PyMuPDF
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from app.config import settings


# Supported file extensions
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt"}


def get_file_extension(filename: str) -> str:
    """Get normalized file extension."""
    return os.path.splitext(filename)[1].lower()


def is_supported(filename: str) -> bool:
    """Check if a file type is supported."""
    return get_file_extension(filename) in SUPPORTED_EXTENSIONS


# ── PDF ──────────────────────────────────────────────────────────

def extract_text_from_pdf(file_path: str) -> list[dict]:
    """Extract text from a PDF file, page by page."""
    pages = []
    doc = fitz.open(file_path)

    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text("text")

        if text.strip():
            pages.append({
                "page_num": page_num + 1,
                "text": text.strip()
            })

    doc.close()
    return pages


def get_pdf_page_count(file_path: str) -> int:
    """Get the total number of pages in a PDF."""
    doc = fitz.open(file_path)
    count = len(doc)
    doc.close()
    return count


# ── DOCX ─────────────────────────────────────────────────────────

def extract_text_from_docx(file_path: str) -> list[dict]:
    """
    Extract text from a DOCX file.

    Treats each paragraph as part of a virtual 'page'.
    Groups every ~3000 characters into a virtual page for chunking.
    """
    from docx import Document as DocxDocument

    doc = DocxDocument(file_path)
    full_text = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            full_text.append(text)

    # Group paragraphs into virtual pages (~3000 chars each)
    pages = []
    current_text = ""
    page_num = 1

    for para_text in full_text:
        current_text += para_text + "\n\n"
        if len(current_text) >= 3000:
            pages.append({"page_num": page_num, "text": current_text.strip()})
            current_text = ""
            page_num += 1

    if current_text.strip():
        pages.append({"page_num": page_num, "text": current_text.strip()})

    return pages


# ── PPTX ─────────────────────────────────────────────────────────

def extract_text_from_pptx(file_path: str) -> list[dict]:
    """
    Extract text from a PPTX file.

    Each slide is treated as a separate page.
    """
    from pptx import Presentation

    prs = Presentation(file_path)
    pages = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        texts.append(text)

        slide_text = "\n".join(texts)
        if slide_text.strip():
            pages.append({
                "page_num": slide_num,
                "text": slide_text.strip()
            })

    return pages


# ── TXT ──────────────────────────────────────────────────────────

def extract_text_from_txt(file_path: str) -> list[dict]:
    """
    Extract text from a plain text file.

    Groups every ~3000 characters into a virtual page.
    """
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()

    if not content.strip():
        return []

    # Split into virtual pages
    pages = []
    page_num = 1
    # Split by double newlines first for natural breaks
    paragraphs = content.split("\n\n")
    current_text = ""

    for para in paragraphs:
        current_text += para.strip() + "\n\n"
        if len(current_text) >= 3000:
            pages.append({"page_num": page_num, "text": current_text.strip()})
            current_text = ""
            page_num += 1

    if current_text.strip():
        pages.append({"page_num": page_num, "text": current_text.strip()})

    return pages


# ── Unified entry point ──────────────────────────────────────────

def extract_text(file_path: str, filename: str) -> list[dict]:
    """
    Extract text from any supported document format.

    Returns list of dicts with 'page_num' and 'text' keys.
    """
    ext = get_file_extension(filename)

    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext == ".pptx":
        return extract_text_from_pptx(file_path)
    elif ext == ".txt":
        return extract_text_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def get_page_count(file_path: str, filename: str) -> int:
    """Get the page/slide count for a document."""
    ext = get_file_extension(filename)

    if ext == ".pdf":
        return get_pdf_page_count(file_path)
    else:
        # For non-PDF, count the virtual pages from extraction
        pages = extract_text(file_path, filename)
        return len(pages)


# ── Chunking (unchanged) ────────────────────────────────────────

def chunk_pages(
    pages: list[dict],
    doc_id: str,
    doc_name: str,
    user_id: str = None,
) -> list[Document]:
    """
    Split extracted pages into smaller chunks for embedding.

    Each chunk preserves metadata about its source document and page number,
    which is critical for source citations.
    """
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
        length_function=len,
    )

    all_chunks = []

    for page in pages:
        metadata = {
            "doc_id": doc_id,
            "doc_name": doc_name,
            "page_num": page["page_num"],
        }
        if user_id:
            metadata["user_id"] = user_id

        chunks = splitter.create_documents(
            texts=[page["text"]],
            metadatas=[metadata]
        )
        all_chunks.extend(chunks)

    return all_chunks
